# -*- coding: utf-8 -*-
"""
Build script: Assembles all modules, outputs build_master_deck.js, runs node to generate PPTX, validates OpenXML, and prints verification summary.
"""

import sys
import os
import subprocess
import shutil

# Import modules
from theme import JS_HEADER
from module0 import SLIDES_0
from module1 import SLIDES_1
from module2 import SLIDES_2
from module3 import SLIDES_3
from module4 import SLIDES_4
from module5 import SLIDES_5

TARGET_JS = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'build_master_deck.js'))
TARGET_PPTX_NAME = "Anthropic 2 万亿 IPO 的背后：编码模型与 AGI_Gemini.pptx"
TARGET_PPTX_ALT = "Anthropic_2万亿IPO的背后_编码模型与AGI_Gemini.pptx"

def main():
    print("=== Assembling Master 45-Slide Generator ===")
    
    js_footer = f"""
// Save Presentation
const outputFile = '{TARGET_PPTX_NAME}';
pres.writeFile({{ fileName: outputFile }})
  .then(f => {{
    console.log('Successfully written master deck:', outputFile);
  }})
  .catch(err => {{
    console.error('Error generating presentation:', err);
    process.exit(1);
  }});
"""

    full_js = "\n".join([
        JS_HEADER,
        SLIDES_0,
        SLIDES_1,
        SLIDES_2,
        SLIDES_3,
        SLIDES_4,
        SLIDES_5,
        js_footer
    ])

    with open(TARGET_JS, 'w', encoding='utf-8') as f:
        f.write(full_js)
    
    print(f"Written JavaScript generator to: {TARGET_JS} ({len(full_js)} bytes)")

    # Run node
    cwd = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    print("Running node build_master_deck.js...")
    res = subprocess.run(["node", "build_master_deck.js"], cwd=cwd, capture_output=True, text=True)
    print("Node output:", res.stdout)
    if res.returncode != 0:
        print("Node error:", res.stderr)
        sys.exit(1)

    # Copy to alternate file name
    pptx_path = os.path.join(cwd, TARGET_PPTX_NAME)
    pptx_alt_path = os.path.join(cwd, TARGET_PPTX_ALT)
    if os.path.exists(pptx_path):
        shutil.copyfile(pptx_path, pptx_alt_path)
        print(f"Synced copy to {TARGET_PPTX_ALT}")

    # Validate with official script
    val_script = "/Users/tangpeng/.gemini/config/skills/pptx/scripts/office/validate.py"
    if os.path.exists(val_script):
        print("Running official OpenXML validator...")
        val_res = subprocess.run(["python3", val_script, pptx_path], capture_output=True, text=True)
        print("Validator output:", val_res.stdout)
        if val_res.returncode != 0:
            print("Validator stderr:", val_res.stderr)
            sys.exit(1)

    # Verify slide count using python-pptx
    import pptx
    prs = pptx.Presentation(pptx_path)
    slide_count = len(prs.slides)
    print(f"\n==========================================")
    print(f"VERIFICATION SUCCESSFUL: {slide_count} SLIDES GENERATED")
    print(f"==========================================")
    for idx, slide in enumerate(prs.slides):
        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text[:40] + "..."
        # Extract title if present
        title_text = "No Title"
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                lines = [l.strip() for l in shape.text_frame.text.split('\n') if l.strip()]
                if lines:
                    title_text = lines[0][:45]
                    break
        print(f"Slide {idx+1:02d}: {title_text:<45} | Notes: {notes_text}")

if __name__ == '__main__':
    main()
